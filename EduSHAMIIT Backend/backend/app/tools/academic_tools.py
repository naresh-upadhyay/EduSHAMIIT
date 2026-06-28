import os
import re
from langchain_core.tools import tool
from app.services.supabase_client import get_supabase
from app.middleware.auth import get_current_user_id

# ─── Unicode maps and math cleaner ───────────────────────────────────────────

SUPERSCRIPT_MAP = {
    '0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
    '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
    '+': '⁺', '-': '⁻', '=': '⁼', '(': '⁽', ')': '⁾',
    'n': 'ⁿ', 'x': 'ˣ', 'i': 'ⁱ', 'r': 'ʳ', 't': 'ᵗ',
    'a': 'ᵃ', 'b': 'ᵇ', 'c': 'ᶜ', 'd': 'ᵈ', 'e': 'ᵉ',
    'f': 'ᶠ', 'g': 'ᵍ', 'h': 'ʰ', 'j': 'ʲ', 'k': 'ᵏ',
    'l': 'ˡ', 'm': 'ᵐ', 'o': 'ᵒ', 'p': 'ᵖ', 's': 'ˢ',
    'u': 'ᵘ', 'v': 'ᵛ', 'w': 'ʷ', 'y': 'ʸ', 'z': 'ᶻ',
}

SUBSCRIPT_MAP = {
    '0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄',
    '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉',
    '+': '₊', '-': '₋', '=': '₌', '(': '₍', ')': '₎',
    'a': 'ₐ', 'e': 'ₑ', 'i': 'ᵢ', 'j': 'ⱼ', 'k': 'ₖ',
    'l': 'ₗ', 'm': 'ₘ', 'n': 'ₙ', 'o': 'ₒ', 'p': 'ₚ',
    'r': 'ᵣ', 's': 'ₛ', 't': 'ₜ', 'u': 'ᵤ', 'v': 'ᵥ',
    'x': 'ₓ',
}

LATEX_MAP = {
    # Operators
    r'\cdot':       '·',
    r'\times':      '×',
    r'\div':        '÷',
    r'\pm':         '±',
    r'\mp':         '∓',
    r'\approx':     '≈',
    r'\neq':        '≠',
    r'\ne':         '≠',
    r'\leq':        '≤',
    r'\geq':        '≥',
    r'\le':         '≤',
    r'\ge':         '≥',
    r'\ll':         '≪',
    r'\gg':         '≫',
    r'\sim':        '~',
    r'\simeq':      '≃',
    r'\equiv':      '≡',
    r'\propto':     '∝',
    r'\in':         '∈',
    r'\notin':      '∉',
    r'\subset':     '⊂',
    r'\supset':     '⊃',
    r'\cup':        '∪',
    r'\cap':        '∩',
    r'\emptyset':   '∅',
    # Arrows
    r'\rightarrow': '→',
    r'\leftarrow':  '←',
    r'\Rightarrow': '⇒',
    r'\Leftarrow':  '⇐',
    r'\leftrightarrow': '↔',
    r'\Leftrightarrow': '⟺',
    r'\leftharpoons': '⇌',
    r'\rightleftharpoons': '⇌',
    r'\to':         '→',
    r'\gets':       '←',
    r'\uparrow':    '↑',
    r'\downarrow':  '↓',
    # Dots & misc
    r'\ldots':      '…',
    r'\cdots':      '···',
    r'\vdots':      '⋮',
    r'\ddots':      '⋱',
    r'\therefore':  '∴',
    r'\because':    '∵',
    r'\forall':     '∀',
    r'\exists':     '∃',
    # Geometry / trig
    r'\angle':      '∠',
    r'\perp':       '⊥',
    r'\parallel':   '∥',
    r'\circ':       '°',
    r'\degree':     '°',
    r'^\circ':      '°',
    r'^\degree':    '°',
    r'^o':          '°',
    # Calc / analysis
    r'\int':        '∫',
    r'\iint':       '∬',
    r'\iiint':      '∭',
    r'\oint':       '∮',
    r'\sum':        '∑',
    r'\prod':       '∏',
    r'\partial':    '∂',
    r'\nabla':      '∇',
    r'\infty':      '∞',
    r'\sqrt':       '√',
    # Named functions
    r'\sin':  'sin',
    r'\cos':  'cos',
    r'\tan':  'tan',
    r'\cot':  'cot',
    r'\sec':  'sec',
    r'\csc':  'csc',
    r'\arcsin': 'arcsin',
    r'\arccos': 'arccos',
    r'\arctan': 'arctan',
    r'\sinh': 'sinh',
    r'\cosh': 'cosh',
    r'\tanh': 'tanh',
    r'\log':  'log',
    r'\ln':   'ln',
    r'\exp':  'exp',
    r'\lim':  'lim',
    r'\max':  'max',
    r'\min':  'min',
    r'\gcd':  'gcd',
    r'\lcm':  'lcm',
    r'\det':  'det',
    r'\dim':  'dim',
    r'\ker':  'ker',
    # Lowercase Greek
    r'\alpha':   'α',
    r'\beta':    'β',
    r'\gamma':   'γ',
    r'\delta':   'δ',
    r'\epsilon': 'ε',
    r'\varepsilon': 'ε',
    r'\zeta':    'ζ',
    r'\eta':     'η',
    r'\theta':   'θ',
    r'\vartheta':'ϑ',
    r'\iota':    'ι',
    r'\kappa':   'κ',
    r'\lambda':  'λ',
    r'\mu':      'μ',
    r'\nu':      'ν',
    r'\xi':      'ξ',
    r'\pi':      'π',
    r'\varpi':   'ϖ',
    r'\rho':     'ρ',
    r'\varrho':  'ϱ',
    r'\sigma':   'σ',
    r'\varsigma':'ς',
    r'\tau':     'τ',
    r'\upsilon': 'υ',
    r'\phi':     'φ',
    r'\varphi':  'φ',
    r'\chi':     'χ',
    r'\psi':     'ψ',
    r'\omega':   'ω',
    # Uppercase Greek
    r'\Gamma':   'Γ',
    r'\Delta':   'Δ',
    r'\Theta':   'Θ',
    r'\Lambda':  'Λ',
    r'\Xi':      'Ξ',
    r'\Pi':      'Π',
    r'\Sigma':   'Σ',
    r'\Upsilon': 'Υ',
    r'\Phi':     'Φ',
    r'\Psi':     'Ψ',
    r'\Omega':   'Ω',
    # Brackets
    r'\lfloor': '⌊', r'\rfloor': '⌋',
    r'\lfloor ': '⌊', r'\rfloor ': '⌋',
    r'\lceil':  '⌈', r'\rceil':  '⌉',
    r'\langle': '⟨', r'\rangle': '⟩',
    # Font wrappers (remove, keep content)
    r'\text':    '',
    r'\mathrm':  '',
    r'\mathbf':  '',
    r'\mathit':  '',
    r'\mathbb':  '',
    r'\boldsymbol': '',
    r'\displaystyle': '',
    r'\textstyle': '',
    # Spacing
    r'\,': ' ', r'\;': ' ', r'\:': ' ', r'\!': '',
    r'\quad': '  ', r'\qquad': '   ',
    # Brackets \left / \right
    r'\left(':  '(',  r'\right)': ')',
    r'\left[':  '[',  r'\right]': ']',
    r'\left\{': '{',  r'\right\}': '}',
    r'\left|':  '|',  r'\right|': '|',
    r'\left':   '',   r'\right':  '',
    # Misc
    r'\limits': '',
    r'\bullet': '•',
    r'\star':   '★',
    r'\dagger': '†',
    r'\ddagger': '‡',
    r'\hbar':   'ℏ',
    r'\ell':    'ℓ',
    r'\Re':     'ℜ',
    r'\Im':     'ℑ',
    r'\aleph':  'ℵ',
}

def clean_math_to_unicode(text: str) -> str:
    if not text or not isinstance(text, str):
        return text

    # 1. Superscript / Subscript character builders
    def to_superscript(s: str) -> str:
        return "".join(SUPERSCRIPT_MAP.get(c, c) for c in s)

    def to_subscript(s: str) -> str:
        return "".join(SUBSCRIPT_MAP.get(c, c) for c in s)

    # 2. Local cleaner for individual math blocks (delimiters already stripped)
    def clean_single_block(math: str) -> str:
        r = math.strip()
        # wrappers \text{...} etc -> extract content
        r = re.sub(r'\\(?:text|mathrm|mathbf|mathit|mathbb|boldsymbol)\{([^{}]*)\}', r'\1', r)

        # \sqrt{x} -> √(x), \sqrt[n]{x} -> ⁿ√(x)
        r = re.sub(r'\\sqrt\[([^\]]+)\]\{([^{}]*)\}', lambda m: f"{to_superscript(m.group(1))}√({m.group(2)})", r)
        r = re.sub(r'\\sqrt\{([^{}]*)\}', r'√(\1)', r)
        r = r.replace(r'\sqrt', '√')

        # \frac{num}{den} -> num/den
        frac_rx = re.compile(r'\\frac\{([^{}]*)\}\{([^{}]*)\}')
        for _ in range(4):
            if not frac_rx.search(r):
                break
            r = frac_rx.sub(lambda m: f"{m.group(1)}/{m.group(2)}", r)

        # Apply symbol map
        keys = sorted(LATEX_MAP.keys(), key=len, reverse=True)
        for key in keys:
            r = r.replace(key, LATEX_MAP[key])

        # Remove remaining unknown \cmd{...}
        r = re.sub(r'\\[a-zA-Z]+\{([^{}]*)\}', r'\1', r)
        r = re.sub(r'\\[a-zA-Z]+', '', r)

        # Superscripts ^{...} and ^x
        r = re.sub(r'\^\{([^{}]*)\}', lambda m: to_superscript(m.group(1)), r)
        r = re.sub(r'\^([0-9a-zA-Z+\-])', lambda m: to_superscript(m.group(1)), r)

        # Subscripts _{...} and _x
        r = re.sub(r'_\{([^{}]*)\}', lambda m: to_subscript(m.group(1)), r)
        r = re.sub(r'_([0-9a-zA-Z+\-])', lambda m: to_subscript(m.group(1)), r)

        r = r.replace('{', '').replace('}', '')
        r = r.replace('\\\\', ' ')
        return re.sub(r'  +', ' ', r).strip()

    # 3. Main processing loop
    r = text

    # Display math blocks $$...$$ and \[...\]
    r = re.compile(r'\$\$(.+?)\$\$', re.DOTALL).sub(lambda m: f"\n{clean_single_block(m.group(1))}\n", r)
    r = re.compile(r'\\\[(.+?)\\\]', re.DOTALL).sub(lambda m: f"\n{clean_single_block(m.group(1))}\n", r)

    # Inline math blocks $...$ and \(...\)
    r = re.sub(r'\$([^\$\n]+)\$', lambda m: clean_single_block(m.group(1)), r)
    r = re.compile(r'\\\((.+?)\\\)', re.DOTALL).sub(lambda m: clean_single_block(m.group(1)), r)

    # Bare symbols outside delimiters
    keys = sorted(LATEX_MAP.keys(), key=len, reverse=True)
    for key in keys:
        if not LATEX_MAP[key]:
            continue
        r = r.replace(key, LATEX_MAP[key])

    # Bare \frac outside delimiters
    frac_rx = re.compile(r'\\frac\{([^{}]*)\}\{([^{}]*)\}')
    for _ in range(4):
        if not frac_rx.search(r):
            break
        r = frac_rx.sub(lambda m: f"{m.group(1)}/{m.group(2)}", r)

    # Bare \sqrt outside delimiters
    r = re.sub(r'√\{([^{}]*)\}', r'√(\1)', r)
    r = re.sub(r'\\[a-zA-Z]+\{([^{}]*)\}', r'\1', r)
    r = re.sub(r'\\[a-zA-Z]+', '', r)

    # Bare superscripts
    r = re.sub(r'(?<=\w)\^\{([^{}]+)\}', lambda m: to_superscript(m.group(1)), r)
    r = re.sub(r'(?<=\w)\^([0-9+\-])', lambda m: to_superscript(m.group(1)), r)

    # Bare subscripts
    r = re.sub(r'(?<=[A-Za-z])_\{([^{}]+)\}', lambda m: to_subscript(m.group(1)), r)
    r = re.sub(r'(?<=[A-Za-z])_([0-9])', lambda m: to_subscript(m.group(1)), r)

    # Chemical subscripts H2O -> H₂O
    common_chem = {
      'H', 'He', 'Li', 'Be', 'B', 'C', 'N', 'O', 'F', 'Ne',
      'Na', 'Mg', 'Al', 'Si', 'P', 'S', 'Cl', 'Ar', 'K', 'Ca',
      'Sc', 'Ti', 'V', 'Cr', 'Mn', 'Fe', 'Co', 'Ni', 'Cu', 'Zn',
      'Ga', 'Ge', 'As', 'Se', 'Br', 'Kr', 'Rb', 'Sr', 'Y', 'Zr',
      'Ag', 'Cd', 'In', 'Sn', 'Sb', 'Te', 'I', 'Xe',
      'Ba', 'La', 'Ce', 'W', 'Re', 'Os', 'Ir', 'Pt', 'Au', 'Hg',
      'Tl', 'Pb', 'Bi', 'Ra', 'U', 'Pu'
    }
    r = re.sub(r'\b([A-Z][a-z]?)([0-9]+)(?=[A-Z0-9])|\b([A-Z][a-z]?)([0-9]+)\b',
               lambda m: f"{m.group(1) or m.group(3)}{to_subscript(m.group(2) or m.group(4))}" if (m.group(1) or m.group(3)) in common_chem else m.group(0),
               r)

    # Stray dollar signs
    r = r.replace(r'\$', '\u0024')
    r = r.replace('$', '')

    # Leftover bracket wrappers
    r = r.replace(r'\(', '').replace(r'\)', '').replace(r'\[', '').replace(r'\]', '')

    # Caret-degree / caret-circle patterns
    r = r.replace('^°', '°').replace('^∘', '°').replace('^o', '°')

    # Clean up escaped single quotes
    r = r.replace(r"\'", "'")
    
    # Clean up backslash space and stray backslashes
    r = re.sub(r'\\([^\w\s])', r'\1', r)
    r = r.replace(r'\ ', ' ')
    r = r.replace('\\', '')

    return r




def get_academic_tools(school_id: str):
    """Return academic AI tools scoped to this school."""

    def _llm():
        from langchain_google_genai import ChatGoogleGenerativeAI
        return ChatGoogleGenerativeAI(
            model=os.getenv("GEMINI_MODEL", "gemini-1.5-flash"),
            temperature=0.7,
            google_api_key=os.getenv("GOOGLE_API_KEY", ""),
        )

    @tool
    def generate_questions(
        subject: str,
        topic: str,
        difficulty: str = "medium",
        num_questions: int = 5,
        question_type: str = "mcq",
    ) -> dict:
        """Generate exam or practice questions for a subject and topic using AI.
        Input: subject name, topic, difficulty (easy/medium/hard),
        number of questions (default 5), question type (mcq/short_answer/long_answer).
        Use when a teacher asks to create questions, make a quiz, or generate test paper."""
        try:
            llm = _llm()
            from langchain_core.messages import HumanMessage
            q_type = question_type.replace("_", " ")
            prompt = (
                f"Generate {num_questions} {q_type} questions for {subject} "
                f"on the topic \"{topic}\" at {difficulty} difficulty level. "
                "For each question provide: question text, options (if MCQ, 4 choices), "
                "the correct answer, and a brief explanation. "
                "Format as a clean numbered list. NCERT-aligned."
            )
            response = llm.invoke([HumanMessage(content=prompt)])
            return {
                "success": True,
                "subject": subject,
                "topic": topic,
                "difficulty": difficulty,
                "question_type": question_type,
                "questions": response.content,
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    @tool
    def create_lesson_plan(
        subject: str,
        topic: str,
        grade: str,
        duration_minutes: int = 40,
        learning_objectives: str = "",
    ) -> dict:
        """Create a detailed lesson plan for a teacher using AI.
        Input: subject, topic, grade/class, duration in minutes (default 40),
        optional learning objectives.
        Use when a teacher asks for a lesson plan, teaching plan, or class preparation."""
        try:
            llm = _llm()
            from langchain_core.messages import HumanMessage
            obj = f" with learning objectives: {learning_objectives}" if learning_objectives else ""
            prompt = (
                f"Create a detailed {duration_minutes}-minute lesson plan for {subject} "
                f"on the topic \"{topic}\" for Grade {grade}{obj}. "
                "Include: Learning Objectives, Materials Needed, Lesson Structure "
                "with time allocation, Activities, Differentiation strategies, "
                "Assessment methods, and Homework. "
                "Make it practical and NCERT-aligned."
            )
            response = llm.invoke([HumanMessage(content=prompt)])
            return {
                "success": True,
                "subject": subject,
                "topic": topic,
                "grade": grade,
                "duration_minutes": duration_minutes,
                "lesson_plan": response.content,
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    @tool
    def get_timetable(class_name: str = "", teacher_id: str = "") -> dict:
        """Get the school timetable for a class or teacher.
        Input: class_name (e.g. '10A') or teacher_id (UUID), or both.
        Use when a teacher or admin asks about the class schedule or teacher timetable."""
        try:
            supabase = get_supabase()
            query = supabase.table("timetable").select("*").eq("school_id", school_id)
            if class_name:
                query = query.eq("class_name", class_name)
            if teacher_id:
                query = query.eq("teacher_id", teacher_id)
            resp = query.order("day").order("period").execute()
            entries = resp.data or []
            timetable: dict = {}
            for entry in entries:
                day = entry.get("day", "Unknown")
                if day not in timetable:
                    timetable[day] = []
                timetable[day].append({
                    "period": entry.get("period"),
                    "subject": entry.get("subject"),
                    "teacher": entry.get("teacher_name", ""),
                    "room": entry.get("room", ""),
                    "start_time": entry.get("start_time", ""),
                    "end_time": entry.get("end_time", ""),
                })
            return {
                "success": True,
                "class_name": class_name,
                "teacher_id": teacher_id,
                "timetable": timetable,
                "total_entries": len(entries),
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    @tool
    def generate_image(prompt: str) -> dict:
        """Generate a beautiful, premium image using AI based on a description prompt.
        Input: a clear description of the image to generate.
        Use when the user asks to draw, generate, show, create, or display an image, picture, or illustration in chat."""
        import urllib.parse
        encoded = urllib.parse.quote(prompt.strip())
        url = f"https://image.pollinations.ai/prompt/{encoded}?width=1024&height=1024&nologo=true&enhance=false&model=turbo"
        markdown_link = f"![Generated Image]({url})"
        return {
            "success": True,
            "prompt": prompt,
            "image_url": url,
            "markdown": markdown_link,
            "result": f"Here is your generated image:\n\n{markdown_link}\n\n*Prompt: {prompt}*"
        }

    @tool
    def generate_document(doc_type: str, filename: str, content: str) -> dict:
        """Generate a document (PDF, Excel spreadsheet, or CSV/text file) for download.
        Input:
          - doc_type: 'pdf' or 'xlsx' or 'csv'
          - filename: preferred name of the file (e.g. 'lesson_plan.pdf', 'grades.xlsx')
          - content: structured text or formatted data:
            - For PDF: formatted Markdown text. Can include # headers, **bold**, *italic*,
              bullet lists (- item), numbered lists (1. item), tables (| col |), code blocks.
              Do NOT use LaTeX math ($...$); write math as plain text instead.
            - For Excel (xlsx) / CSV: structured JSON string of rows and columns,
              e.g. '[[col1, col2], [val1, val2]]'.
        Use when a user asks to generate, export, download, or create a PDF, Excel spreadsheet, or CSV document."""
        import os as _os
        ext = doc_type.lower().strip()
        if not filename.endswith(f".{ext}"):
            filename = f"{filename}.{ext}"

        assets_dir = _os.path.join(_os.getcwd(), "assets")
        _os.makedirs(assets_dir, exist_ok=True)
        file_path = _os.path.join(assets_dir, filename)

        try:
            # ──────────────────────────────────────────────────────────────────
            #  PDF Generation – Professional, Beautiful Output
            # ──────────────────────────────────────────────────────────────────
            if ext == "pdf":
                content = clean_math_to_unicode(content)
                import re
                from reportlab.lib.pagesizes import letter
                from reportlab.lib import colors
                from reportlab.lib.units import inch
                from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
                from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
                from reportlab.platypus import (
                    SimpleDocTemplate, Paragraph, Spacer, HRFlowable,
                    Table, TableStyle, KeepTogether
                )

                # ── Colour palette ────────────────────────────────────────────
                C_BRAND      = colors.HexColor("#4F46E5")
                C_BRAND_LIGHT= colors.HexColor("#EEF2FF")
                C_H2         = colors.HexColor("#1E293B")
                C_H3         = colors.HexColor("#334155")
                C_BODY       = colors.HexColor("#1E293B")
                C_MUTED      = colors.HexColor("#64748B")
                C_CODE_BG    = colors.HexColor("#F1F5F9")
                C_ROW_ODD    = colors.HexColor("#F8FAFC")
                C_TABLE_LINE = colors.HexColor("#CBD5E1")
                C_RULE       = colors.HexColor("#E2E8F0")

                # ── LaTeX / math cleaner ──────────────────────────────────────
                LATEX_MAP = {
                    r'\frac': '', r'\cdot': '·', r'\times': '×', r'\div': '÷',
                    r'\pm': '±', r'\approx': '≈', r'\neq': '≠',
                    r'\leq': '≤', r'\geq': '≥', r'\le': '≤', r'\ge': '≥',
                    r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ', r'\delta': 'δ',
                    r'\epsilon': 'ε', r'\theta': 'θ', r'\lambda': 'λ', r'\mu': 'μ',
                    r'\pi': 'π', r'\sigma': 'σ', r'\omega': 'ω',
                    r'\Delta': 'Δ', r'\Sigma': 'Σ', r'\Omega': 'Ω', r'\Pi': 'Π',
                    r'\infty': '∞', r'\int': '∫', r'\sum': 'Σ', r'\prod': 'Π',
                    r'\partial': '∂', r'\nabla': '∇',
                    r'\rightarrow': '→', r'\leftarrow': '←',
                    r'\Rightarrow': '⇒', r'\Leftarrow': '⇐', r'\to': '→',
                    r'\ldots': '…', r'\cdots': '···',
                    r'\therefore': '∴', r'\because': '∵',
                    r'\angle': '∠', r'\perp': '⊥', r'\parallel': '∥',
                    r'\sqrt': '√', r'\sin': 'sin', r'\cos': 'cos', r'\tan': 'tan',
                    r'\log': 'log', r'\ln': 'ln', r'\lim': 'lim',
                    r'\text': '', r'\mathrm': '', r'\mathbf': '', r'\mathit': '',
                }

                def _latex_expr(expr: str) -> str:
                    # \frac{a}{b} → a/b
                    expr = re.sub(r'\\frac\{([^}]*)\}\{([^}]*)\}', r'\1/\2', expr)
                    # \sqrt{a} → √(a)
                    expr = re.sub(r'\\sqrt\{([^}]*)\}', r'√(\1)', expr)
                    # \cmd{text} → text
                    expr = re.sub(r'\\[a-zA-Z]+\{([^}]*)\}', r'\1', expr)
                    # subscripts _{x} → _x
                    expr = re.sub(r'_\{([^}]*)\}', r'_\1', expr)
                    # superscripts ^{x} → ^x
                    expr = re.sub(r'\^\{([^}]*)\}', r'^\1', expr)
                    for lat, uni in LATEX_MAP.items():
                        expr = expr.replace(lat, uni)
                    expr = re.sub(r'\\[a-zA-Z]+', '', expr)
                    return expr.replace('{', '').replace('}', '').strip()

                def clean_math(text: str) -> str:
                    """Remove LaTeX $...$ and $$...$$ and replace with readable text."""
                    text = re.sub(r'\$\$(.+?)\$\$', lambda m: _latex_expr(m.group(1)), text, flags=re.DOTALL)
                    text = re.sub(r'\$(.+?)\$', lambda m: _latex_expr(m.group(1)), text)
                    return text

                def md_to_rl(text: str) -> str:
                    """Convert markdown line to ReportLab XML markup."""
                    # ── Step 1: clean math/LaTeX expressions ─────────────────
                    text = clean_math(text)
                    
                    # ── Step 2: clean up backslash artifacts ─────────────────
                    text = text.replace(r"\'", "'")
                    text = re.sub(r'\\([^\w\s])', r'\1', text)
                    text = text.replace(r'\ ', ' ')
                    text = text.replace('\\', '')
                    
                    # ── Step 3: Unicode sub/superscript → plain text BEFORE XML ──
                    # (must happen before XML-escaping so we don't corrupt tag chars)
                    rev_sub = {v: k for k, v in SUBSCRIPT_MAP.items()}
                    rev_super = {v: k for k, v in SUPERSCRIPT_MAP.items()}
                    sub_chars = re.escape("".join(rev_sub.keys()))
                    super_chars = re.escape("".join(rev_super.keys()))
                    if rev_sub:
                        text = re.sub(f"([{sub_chars}]+)",
                            lambda m: "[SUB]" + "".join(rev_sub.get(c, c) for c in m.group(1)) + "[/SUB]",
                            text)
                    if rev_super:
                        text = re.sub(f"([{super_chars}]+)",
                            lambda m: "[SUP]" + "".join(rev_super.get(c, c) for c in m.group(1)) + "[/SUP]",
                            text)
                    
                    # ── Step 4: chemical subscripts and caret-superscripts ────
                    # Only match digits NOT already wrapped in [SUB]/[SUP] placeholders
                    # Apply before XML escaping
                    text = re.sub(r'([A-Za-z])([0-9]+)(?=[A-Z\s,.)!]|$)', r'\1[SUB]\2[/SUB]', text)
                    text = re.sub(r'(\w)\^([A-Za-z0-9+\-]+)', r'\1[SUP]\2[/SUP]', text)
                    text = re.sub(r'([A-Za-z])_([0-9]+)', r'\1[SUB]\2[/SUB]', text)

                    # ── Step 5: XML-escape special chars BEFORE inserting real tags ──
                    text = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    
                    # ── Step 6: Now convert [SUB]/[SUP] placeholders to real XML tags ──
                    text = text.replace('[SUB]', '<sub>').replace('[/SUB]', '</sub>')
                    text = text.replace('[SUP]', '<sup>').replace('[/SUP]', '</sup>')

                    # ── Step 7: Inline markdown (bold, italic, code) ──────────
                    text = re.sub(r'\*\*\*(.*?)\*\*\*', r'<b><i>\1</i></b>', text)
                    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
                    text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
                    text = re.sub(r'_(.*?)_', r'<i>\1</i>', text)
                    text = re.sub(r'`([^`]+)`', r'<font face="Courier" size="9">\1</font>', text)

                    # ── Step 8: Collapse accidental nested sub/sup tags ───────
                    # e.g. <sub><sub>2</sub></sub> → <sub>2</sub>
                    for _ in range(3):
                        text = re.sub(r'<(sub|sup)><\1>([^<]*)</\1></\1>', r'<\1>\2</\1>', text)

                    return text

                # ── Style sheet ───────────────────────────────────────────────
                styles = getSampleStyleSheet()

                sty_h1 = ParagraphStyle(
                    'H1', parent=styles['Normal'],
                    fontSize=24, leading=30, fontName='Helvetica-Bold',
                    textColor=C_BRAND, spaceBefore=0, spaceAfter=4,
                    alignment=TA_LEFT,
                )
                sty_h2 = ParagraphStyle(
                    'H2', parent=styles['Normal'],
                    fontSize=16, leading=21, fontName='Helvetica-Bold',
                    textColor=C_H2, spaceBefore=16, spaceAfter=4,
                )
                sty_h3 = ParagraphStyle(
                    'H3', parent=styles['Normal'],
                    fontSize=13, leading=17, fontName='Helvetica-Bold',
                    textColor=C_H3, spaceBefore=10, spaceAfter=3,
                )
                sty_h4 = ParagraphStyle(
                    'H4', parent=styles['Normal'],
                    fontSize=11, leading=15, fontName='Helvetica-BoldOblique',
                    textColor=C_H3, spaceBefore=6, spaceAfter=2,
                )
                sty_body = ParagraphStyle(
                    'Body', parent=styles['Normal'],
                    fontSize=11, leading=16, fontName='Helvetica',
                    textColor=C_BODY, spaceAfter=5, alignment=TA_JUSTIFY,
                )
                sty_bullet = ParagraphStyle(
                    'Bullet', parent=styles['Normal'],
                    fontSize=11, leading=15, fontName='Helvetica',
                    textColor=C_BODY, leftIndent=20, spaceAfter=3,
                )
                sty_num = ParagraphStyle(
                    'Num', parent=styles['Normal'],
                    fontSize=11, leading=15, fontName='Helvetica',
                    textColor=C_BODY, leftIndent=24, firstLineIndent=-24,
                    spaceAfter=3,
                )
                sty_code = ParagraphStyle(
                    'Code', parent=styles['Normal'],
                    fontSize=9, leading=13, fontName='Courier',
                    textColor=C_H2, backColor=C_CODE_BG,
                    leftIndent=10, rightIndent=10,
                    spaceBefore=4, spaceAfter=4,
                    borderPadding=(5, 8, 5, 8),
                )
                sty_quote = ParagraphStyle(
                    'Quote', parent=styles['Normal'],
                    fontSize=11, leading=15, fontName='Helvetica-Oblique',
                    textColor=C_MUTED, leftIndent=18,
                    spaceBefore=3, spaceAfter=3,
                )
                sty_thead = ParagraphStyle(
                    'THead', parent=styles['Normal'],
                    fontSize=10, fontName='Helvetica-Bold',
                    textColor=C_H2, alignment=TA_CENTER,
                )
                sty_tcell = ParagraphStyle(
                    'TCell', parent=styles['Normal'],
                    fontSize=10, fontName='Helvetica',
                    textColor=C_BODY,
                )

                # ── Page header & footer callback ─────────────────────────────
                def _on_page(canvas, doc):
                    canvas.saveState()
                    w, h = letter
                    # Top rule
                    canvas.setStrokeColor(C_BRAND)
                    canvas.setLineWidth(2)
                    canvas.line(0.7*inch, h - 0.45*inch, w - 0.7*inch, h - 0.45*inch)
                    # Footer text
                    canvas.setFont('Helvetica', 8)
                    canvas.setFillColor(C_MUTED)
                    canvas.drawString(0.7*inch, 0.38*inch, "EduSHAMIIT – AI Generated Document")
                    canvas.drawRightString(w - 0.7*inch, 0.38*inch, f"Page {doc.page}")
                    canvas.restoreState()

                # ── Table parser ──────────────────────────────────────────────
                def build_table(rows_raw: list) -> list:
                    """Turn raw markdown table lines into a ReportLab Table."""
                    clean = []
                    for row in rows_raw:
                        if re.match(r'^\s*\|?[-:| ]+\|?\s*$', row):
                            continue
                        cells = [c.strip() for c in re.split(r'\|', row.strip('|')) ]
                        clean.append(cells)
                    if not clean:
                        return []
                    max_c = max(len(r) for r in clean)
                    for r in clean:
                        while len(r) < max_c:
                            r.append('')
                    t_data = []
                    for ridx, row in enumerate(clean):
                        sty = sty_thead if ridx == 0 else sty_tcell
                        t_data.append([Paragraph(md_to_rl(c), sty) for c in row])
                    col_w = (letter[0] - 1.5*inch) / max_c
                    tbl = Table(t_data, colWidths=[col_w]*max_c, repeatRows=1)
                    tbl.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), C_BRAND_LIGHT),
                        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
                         [C_ROW_ODD, colors.white]),
                        ('GRID', (0, 0), (-1, -1), 0.5, C_TABLE_LINE),
                        ('LINEBELOW', (0, 0), (-1, 0), 1.5, C_BRAND),
                        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                        ('FONTSIZE', (0, 0), (-1, -1), 10),
                        ('TOPPADDING', (0, 0), (-1, -1), 5),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                        ('LEFTPADDING', (0, 0), (-1, -1), 8),
                        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ]))
                    return [Spacer(1, 4), tbl, Spacer(1, 8)]

                # ── Markdown → Story ──────────────────────────────────────────
                def md_to_story(md: str) -> list:
                    story = []
                    lines = md.split('\n')
                    i = 0
                    in_code = False
                    code_buf = []
                    in_table = False
                    tbl_buf = []

                    while i < len(lines):
                        line = lines[i]
                        raw = line.strip()

                        # ── code fence ──────────────────────────────────────
                        if raw.startswith('```'):
                            if not in_code:
                                in_code = True
                                code_buf = []
                            else:
                                in_code = False
                                txt = '\n'.join(code_buf)
                                txt = txt.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
                                story.append(Paragraph(txt.replace('\n', '<br/>'), sty_code))
                            i += 1
                            continue

                        if in_code:
                            code_buf.append(line)
                            i += 1
                            continue

                        # ── table row ───────────────────────────────────────
                        if '|' in raw and (raw.startswith('|') or raw.endswith('|')):
                            if not in_table:
                                in_table = True
                                tbl_buf = []
                            tbl_buf.append(raw)
                            i += 1
                            continue
                        else:
                            if in_table:
                                in_table = False
                                story.extend(build_table(tbl_buf))
                                tbl_buf = []

                        # ── blank line ──────────────────────────────────────
                        if not raw:
                            story.append(Spacer(1, 8))
                            i += 1
                            continue

                        # ── horizontal rule ─────────────────────────────────
                        if re.match(r'^(-{3,}|_{3,}|\*{3,})$', raw):
                            story.append(Spacer(1, 4))
                            story.append(HRFlowable(width='100%', thickness=1,
                                                    color=C_RULE, spaceAfter=4))
                            i += 1
                            continue

                        # ── headings ────────────────────────────────────────
                        if re.match(r'^#{1} [^#]', raw):
                            text = raw.lstrip('#').strip()
                            story.append(Spacer(1, 6))
                            story.append(Paragraph(md_to_rl(text), sty_h1))
                            story.append(HRFlowable(width='100%', thickness=2,
                                                    color=C_BRAND, spaceAfter=8))
                            i += 1
                            continue

                        if re.match(r'^#{2} [^#]', raw):
                            text = raw.lstrip('#').strip()
                            story.append(Paragraph(md_to_rl(text), sty_h2))
                            story.append(HRFlowable(width='50%', thickness=1,
                                                    color=C_BRAND, spaceAfter=4))
                            i += 1
                            continue

                        if re.match(r'^#{3} [^#]', raw):
                            text = raw.lstrip('#').strip()
                            story.append(Paragraph(md_to_rl(text), sty_h3))
                            i += 1
                            continue

                        if re.match(r'^#{4,} ', raw):
                            text = raw.lstrip('#').strip()
                            story.append(Paragraph(md_to_rl(text), sty_h4))
                            i += 1
                            continue

                        # ── blockquote ──────────────────────────────────────
                        if raw.startswith('> '):
                            story.append(Paragraph(md_to_rl(raw[2:].strip()), sty_quote))
                            i += 1
                            continue

                        # ── unordered list ───────────────────────────────────
                        if re.match(r'^[-*+] ', raw):
                            text = re.sub(r'^[-*+] ', '', raw)
                            story.append(Paragraph(f'• {md_to_rl(text)}', sty_bullet))
                            i += 1
                            continue

                        # ── ordered list ─────────────────────────────────────
                        m = re.match(r'^(\d+)[.)]\s+(.+)', raw)
                        if m:
                            story.append(
                                Paragraph(f'<b>{m.group(1)}.</b>  {md_to_rl(m.group(2))}',
                                          sty_num)
                            )
                            i += 1
                            continue

                        # ── normal paragraph ─────────────────────────────────
                        story.append(Paragraph(md_to_rl(raw), sty_body))
                        i += 1

                    if in_table and tbl_buf:
                        story.extend(build_table(tbl_buf))

                    return story

                # ── Build document ────────────────────────────────────────────
                doc = SimpleDocTemplate(
                    file_path,
                    pagesize=letter,
                    leftMargin=0.75*inch,
                    rightMargin=0.75*inch,
                    topMargin=0.75*inch,
                    bottomMargin=0.65*inch,
                )
                story = md_to_story(content)
                doc.build(story, onFirstPage=_on_page, onLaterPages=_on_page)

            # ──────────────────────────────────────────────────────────────────
            #  Excel Generation
            # ──────────────────────────────────────────────────────────────────
            elif ext == "xlsx":
                import json
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Sheet1"
                try:
                    data = json.loads(content)
                    if isinstance(data, list):
                        for row_idx, row in enumerate(data, 1):
                            if isinstance(row, list):
                                for col_idx, val in enumerate(row, 1):
                                    cleaned_val = clean_math_to_unicode(str(val))
                                    cell = ws.cell(row=row_idx, column=col_idx, value=cleaned_val)
                                    if row_idx == 1:
                                        cell.font = Font(bold=True, color="FFFFFF")
                                        cell.fill = PatternFill("solid", fgColor="4F46E5")
                                        cell.alignment = Alignment(horizontal="center")
                                    elif row_idx % 2 == 0:
                                        cell.fill = PatternFill("solid", fgColor="F8FAFC")
                    else:
                        ws.cell(row=1, column=1, value=clean_math_to_unicode(content))
                except Exception:
                    for row_idx, line in enumerate(content.split('\n'), 1):
                        for col_idx, col in enumerate(line.split(','), 1):
                            cleaned_col = clean_math_to_unicode(col.strip())
                            ws.cell(row=row_idx, column=col_idx, value=cleaned_col)
                # Auto-size columns
                for col in ws.columns:
                    max_len = max((len(str(cell.value or '')) for cell in col), default=0)
                    ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)
                wb.save(file_path)

            # ──────────────────────────────────────────────────────────────────
            #  CSV Generation
            # ──────────────────────────────────────────────────────────────────
            elif ext == "csv":
                import json
                import csv
                with open(file_path, mode='w', newline='', encoding='utf-8') as f:
                    writer = csv.writer(f)
                    try:
                        data = json.loads(content)
                        if isinstance(data, list):
                            for row in data:
                                if isinstance(row, list):
                                    cleaned_row = [clean_math_to_unicode(str(val)) for val in row]
                                    writer.writerow(cleaned_row)
                                else:
                                    writer.writerow([clean_math_to_unicode(str(row))])
                        else:
                            writer.writerow([clean_math_to_unicode(content)])
                    except Exception:
                        for line in content.split('\n'):
                            cleaned_row = [clean_math_to_unicode(col.strip()) for col in line.split(',')]
                            writer.writerow(cleaned_row)

            # ──────────────────────────────────────────────────────────────────
            #  Plain text fallback
            # ──────────────────────────────────────────────────────────────────
            else:
                cleaned_content = clean_math_to_unicode(content)
                with open(file_path, 'w', encoding='utf-8') as f:
                    f.write(cleaned_content)

            # ── Auto-save to Documents Hub ────────────────────────────────────
            try:
                import uuid
                from app.middleware.auth import get_current_user_id
                user_id = get_current_user_id()
                if user_id:
                    with open(file_path, "rb") as f:
                        file_bytes = f.read()
                    file_size = len(file_bytes)
                    doc_id = str(uuid.uuid4())
                    storage_path = f"documents/{school_id}/{user_id}/{doc_id}.{ext}"
                    from app.config import settings
                    import httpx
                    import mimetypes
                    supabase_url = settings.SUPABASE_URL.rstrip("/")
                    storage_url = f"{supabase_url}/storage/v1/object/{storage_path}"
                    mime = mimetypes.guess_type(filename)[0] or "application/octet-stream"
                    headers = {
                        "apikey": settings.SUPABASE_SERVICE_ROLE_KEY,
                        "Authorization": f"Bearer {settings.SUPABASE_SERVICE_ROLE_KEY}",
                        "Content-Type": mime,
                        "x-upsert": "true",
                    }
                    with httpx.Client(timeout=30.0) as http_client:
                        upload_response = http_client.post(storage_url, headers=headers, content=file_bytes)
                    if upload_response.status_code in (200, 201):
                        from app.middleware.auth import get_public_supabase_url
                        public_url_base = get_public_supabase_url(supabase_url)
                        file_url = f"{public_url_base}/storage/v1/object/public/{storage_path}"
                        from app.services.supabase_client import get_supabase
                        sb = get_supabase()
                        clean_title = filename.rsplit(".", 1)[0].replace("_", " ").title()
                        record = {
                            "id": doc_id,
                            "school_id": school_id,
                            "owner_id": user_id,
                            "title": clean_title,
                            "description": f"AI generated {ext.upper()} document.",
                            "category": "ai_generated",
                            "file_url": file_url,
                            "file_name": filename,
                            "file_size": file_size,
                            "mime_type": mime,
                            "storage_path": storage_path,
                        }
                        sb.table("user_documents").insert(record).execute()
                        print(f"Auto-saved document {filename} for user {user_id}.")
            except Exception as auto_save_err:
                print(f"Auto-save error: {auto_save_err}")

            download_url = f"/api/chat/download/{filename}"
            return {
                "success": True,
                "filename": filename,
                "download_url": download_url,
                "markdown_link": f"[Download {filename}]({download_url})",
                "result": f"I have successfully generated your document! You can download it here: [Download {filename}]({download_url})"
            }
        except Exception as e:
            return {"success": False, "error": str(e)}

    return [generate_questions, create_lesson_plan, get_timetable, generate_image, generate_document]