import json
import uuid
import re
import ast
from typing import AsyncGenerator, Optional, Any
from datetime import datetime

from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.messages import AIMessage, HumanMessage

from app.services.supabase_client import get_supabase
from app.middleware.auth import set_current_user_context
from app.agents.router import TaskType, detect_task, get_llm
from app.agents.prompts import get_system_prompt


def get_gemini_fallbacks(primary_model: str) -> list:
    """Return the list of top 5 available free-tier Gemini models as fallbacks."""
    from langchain_google_genai import ChatGoogleGenerativeAI
    import os
    
    # Ordered list of free-tier Gemini models. High-limit models (Flash) first.
    all_gemini = [
        "gemini-2.5-flash",
        "gemini-2.5-flash-lite",
        "gemini-1.5-flash",
        "gemini-1.5-flash-8b",
        "gemini-2.5-pro",
        "gemini-1.5-pro"
    ]
    fallbacks = []
    
    google_key = os.getenv("GOOGLE_API_KEY")
    if not google_key or google_key.startswith("AIza-placeholder"):
        return []
        
    for m in all_gemini:
        if m != primary_model:
            fallbacks.append(
                ChatGoogleGenerativeAI(
                    model=m,
                    temperature=0.3,
                    max_tokens=8192,
                    streaming=True,
                    google_api_key=google_key
                )
            )
    return fallbacks[:5]


def build_agent(role: str, school_id: str, task_type: str = "qa", user_id: str = None):
    """Build a role-scoped, task-appropriate LangChain agent with fallback LLMs."""
    from langchain.agents import AgentExecutor
    from langchain.agents.format_scratchpad import format_to_tool_messages
    from langchain.agents.output_parsers import ToolsAgentOutputParser
    from langchain_core.runnables import RunnablePassthrough
    from app.tools import get_all_tools, filter_tools_by_role
    from app.agents.router import get_fallback_llms

    llm = get_llm(task_type)
    all_tools = get_all_tools(school_id)
    tools = filter_tools_by_role(all_tools, role, school_id)

    prompt = ChatPromptTemplate.from_messages([
        ("system", get_system_prompt(role, school_id, user_id)),
        MessagesPlaceholder("chat_history"),
        ("human", "{input}\n\nCRITICAL REMINDER: If you are generating, exporting, downloading, or creating any document (PDF, Excel spreadsheet, CSV) or image, you MUST call the appropriate tool ('generate_document' or 'generate_image') first to create the file. NEVER output a download link or markdown link manually or write one you made up; you must only use the exact markdown link returned in the tool's execution result. If you output a link without executing the tool, the file will not exist on the server and the user will get a 404 error."),
        MessagesPlaceholder("agent_scratchpad"),
    ])

    # Bind tools to the primary model
    primary_with_tools = llm.bind_tools(tools)

    # Fetch fallbacks (Gemini fallback list + OpenRouter/GitHub fallback list)
    gemini_fallbacks = get_gemini_fallbacks(getattr(llm, "model", ""))
    other_fallbacks = get_fallback_llms()
    fallbacks = gemini_fallbacks + other_fallbacks
    
    if fallbacks:
        fallbacks_with_tools = [f.bind_tools(tools) for f in fallbacks]
        model_with_tools = primary_with_tools.with_fallbacks(fallbacks_with_tools)
    else:
        model_with_tools = primary_with_tools

    # Construct the tool-calling agent custom pipeline
    agent = (
        RunnablePassthrough.assign(
            agent_scratchpad=lambda x: format_to_tool_messages(
                x["intermediate_steps"]
            )
        )
        | prompt
        | model_with_tools
        | ToolsAgentOutputParser()
    )

    return AgentExecutor(
        agent=agent,
        tools=tools,
        return_intermediate_steps=True,
        max_iterations=5,
        verbose=True,
    )


def load_history(session_id: str) -> list:
    """Load chat history from database."""
    try:
        sb = get_supabase()
        messages = sb.table("ai_chat_history") \
            .select("role, content, tool_calls") \
            .eq("session_id", session_id) \
            .order("created_at", ascending=False) \
            .limit(40).execute()

        history = []
        for msg in reversed(messages.data):
            if msg["role"] == "user":
                history.append(HumanMessage(content=msg["content"]))
            elif msg["role"] == "assistant":
                content = msg["content"]
                tool_calls = msg.get("tool_calls") or {}
                # Clean up hallucinated download links in history
                if "/api/chat/download/" in content and "generate_document" not in tool_calls:
                    content = re.sub(
                        r"\[Download [^\]]+\]\(/api/chat/download/[^\)]+\)", 
                        "(document download link available after generation)", 
                        content
                    )
                history.append(AIMessage(content=content))
        return history
    except Exception:
        return []


async def save_message(session_id: str, user_id: str, role: str, content: str, school_id: str = "", tool_data: dict = None):
    """Save a message to chat history."""
    try:
        sb = get_supabase()
        sb.table("ai_chat_history").insert({
            "school_id": school_id,
            "user_id": user_id,
            "session_id": session_id,
            "role": role,
            "content": content,
            "tool_calls": tool_data,
        }).execute()
    except Exception as e:
        print(f"Save message error: {e}")


def extract_tool_results(result: dict) -> dict:
    """Extract tool results from agent output."""
    tool_data = {}
    steps = result.get("intermediate_steps", [])
    for step in steps:
        if len(step) >= 2:
            action = step[0]
            observation = step[1]
            tool_name = getattr(action, "tool", "unknown")
            tool_data[tool_name] = str(observation)[:500]
    return tool_data


def extract_document_text(doc_bytes: bytes, filename: str) -> str:
    """Extract text content from various document types (PDF, XLSX, CSV, TXT, DOCX)."""
    ext = filename.split(".")[-1].lower()
    
    if ext == "pdf":
        try:
            import io
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(doc_bytes))
            text_parts = []
            for idx, page in enumerate(reader.pages):
                text_parts.append(f"--- Page {idx + 1} ---")
                text_parts.append(page.extract_text() or "")
            return "\n".join(text_parts)
        except Exception as e:
            return f"[Error parsing PDF: {str(e)}]"
            
    elif ext in ["xlsx", "xls"]:
        try:
            import io
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(doc_bytes), read_only=True, data_only=True)
            text_parts = []
            for sheet in wb.sheetnames:
                text_parts.append(f"--- Sheet: {sheet} ---")
                ws = wb[sheet]
                for row in ws.iter_rows(values_only=True):
                    if any(row):
                        text_parts.append(", ".join(str(cell) if cell is not None else "" for cell in row))
            return "\n".join(text_parts)
        except Exception as e:
            return f"[Error parsing Excel: {str(e)}]"
            
    elif ext == "csv":
        try:
            return doc_bytes.decode("utf-8", errors="ignore")
        except Exception as e:
            return f"[Error parsing CSV: {str(e)}]"
            
    elif ext in ["docx", "doc"]:
        try:
            import io
            import zipfile
            import xml.etree.ElementTree as ET
            
            with zipfile.ZipFile(io.BytesIO(doc_bytes)) as docx:
                content_xml = docx.read('word/document.xml')
                root = ET.fromstring(content_xml)
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                paragraphs = []
                for p in root.findall('.//w:p', ns):
                    texts = [t.text for t in p.findall('.//w:t', ns) if t.text]
                    if texts:
                        paragraphs.append("".join(texts))
                return "\n".join(paragraphs)
        except Exception as e:
            try:
                return doc_bytes.decode("utf-8", errors="ignore")
            except Exception:
                return f"[Error parsing Word document: {str(e)}]"
                
    else:
        try:
            return doc_bytes.decode("utf-8", errors="ignore")
        except Exception as e:
            return f"[Unsupported document format: {ext}]"


async def process_message(
    text: str = "",
    image_b64: str = None,
    audio_path: str = None,
    doc_b64: str = None,
    doc_name: str = None,
    user: dict = None,
    session_id: str = "",
    school_id: str = "",
) -> AsyncGenerator[dict, None]:
    """Single entry point for all message types. Yields SSE chunks."""

    # Set user context for tools
    if user:
        set_current_user_context(user)

    # Step 1: Transcribe voice if provided
    if audio_path:
        from app.services.whisper_service import transcribe_file
        text = await transcribe_file(audio_path)
        yield {"type": "text", "content": f"🎤 I heard: {text}\n\n"}

    # Step 2: Handle image directly (skip agent for pure vision tasks)
    if image_b64:
        from app.services.vision_service import analyze_image
        result = await analyze_image(image_b64, text)
        
        # Persist messages in database history
        user_id = user.get("id", "") if user else ""
        user_content = text if text else "Describe this image"
        await save_message(session_id, user_id, "user", f"📷 [Image] {user_content}", school_id)
        await save_message(session_id, user_id, "assistant", result, school_id)
        
        yield {"type": "text", "content": result}
        yield {"type": "done"}
        return

    # Step 2.5: Handle document processing if provided
    original_user_text = text
    if doc_b64 and doc_name:
        import base64
        try:
            doc_bytes = base64.b64decode(doc_b64)
            extracted_text = extract_document_text(doc_bytes, doc_name)
            text = f"[Document Attached: {doc_name}]\n---\n{extracted_text}\n---\n\nUser Question: {original_user_text}"
        except Exception as de:
            text = f"[Error decoding document {doc_name}: {str(de)}]\n\nUser Question: {original_user_text}"

    if not text:
        yield {"type": "text", "content": "Please send a message, voice note, or document."}
        yield {"type": "done"}
        return

    # Step 3: Detect task type and build agent
    task = detect_task(text)
    role = user.get("role", "student") if user else "student"

    try:
        agent = build_agent(role, school_id, task, user_id=user.get("id") if user else None)
        history = load_history(session_id)

        # Step 4: Run agent and stream final answer tokens via astream_events v2
        accumulated_text = ""
        active_tools = {}
        completed_tools = []

        async for event in agent.astream_events(
            {"input": text, "chat_history": history},
            version="v2"
        ):
            kind = event["event"]
            name = event["name"]

            if kind == "on_chat_model_stream":
                data_dict = event.get("data") or {}
                chunk = data_dict.get("chunk")
                if chunk and hasattr(chunk, "content") and chunk.content:
                    token = chunk.content
                    accumulated_text += token
                    yield {"type": "text", "content": token}

            elif kind == "on_tool_start":
                data_dict = event.get("data") or {}
                tool_input_val = data_dict.get("input")
                print(f"DEBUG TOOL START: name={name} input={tool_input_val} type={type(tool_input_val)}")
                active_tools[event["run_id"]] = (name, tool_input_val)
                display_name = name.replace("_", " ").title()
                
                # Show scratchpad thinking / input preview for document generation
                content_preview = ""
                if name == "generate_document" and isinstance(tool_input_val, dict):
                    doc_content = tool_input_val.get("content", "")
                    if doc_content:
                        clean_preview = doc_content.replace("#", "").replace("*", "").strip()
                        preview_text = clean_preview[:120] + "..." if len(clean_preview) > 120 else clean_preview
                        content_preview = f"\n> *Drafting Content: \"{preview_text}\"*\n"
                
                yield {"type": "text", "content": f"\n\n⚙️ *[Executing: {display_name}...]*\n{content_preview}\n"}

            elif kind == "on_tool_end":
                run_id = event["run_id"]
                if run_id in active_tools:
                    tool_name, tool_input = active_tools.pop(run_id)
                    data_dict = event.get("data") or {}
                    tool_output = data_dict.get("output")
                    completed_tools.append((tool_name, tool_input, tool_output))
                    display_name = tool_name.replace("_", " ").title()
                    yield {"type": "text", "content": f"\n*[Completed: {display_name}]* ✅\n\n"}

        # Build final response text and tool results
        output = accumulated_text if accumulated_text else "I couldn't process that request."
        
        # Check if generate_document tool was executed and extract result
        generated_doc = None
        for tool_name, tool_input, tool_output in completed_tools:
            if tool_name == "generate_document":
                download_url = None
                markdown_link = None
                if isinstance(tool_output, dict):
                    download_url = tool_output.get("download_url")
                    markdown_link = tool_output.get("markdown_link")
                elif isinstance(tool_output, str):
                    try:
                        parsed = ast.literal_eval(tool_output)
                        if isinstance(parsed, dict):
                            download_url = parsed.get("download_url")
                            markdown_link = parsed.get("markdown_link")
                    except Exception:
                        m_url = re.search(r"'/api/chat/download/[^']+'", tool_output)
                        if m_url:
                            download_url = m_url.group(0).strip("'")
                        m_link = re.search(r"(\[Download [^\]]+\]\(/api/chat/download/[^\)]+\))", tool_output)
                        if m_link:
                            markdown_link = m_link.group(0)
                
                if download_url and markdown_link:
                    generated_doc = (download_url, markdown_link)
                    break
        
        # Process and safeguard download links in the output text
        original_output = output
        download_pattern = r"\[[^\]]+\]\(/api/chat/download/[^\)]+\)"
        
        if generated_doc:
            download_url, markdown_link = generated_doc
            # Case 1: Document generated successfully. Replace all download links in output with the correct generated link.
            if re.search(download_pattern, output):
                output = re.sub(download_pattern, markdown_link, output)
            else:
                # Append correct link if missing
                output += f"\n\nHere is your generated document: {markdown_link}"
        else:
            # Case 2: Document generation did not occur or failed.
            # Clean up all hallucinated download links to prevent 404 errors.
            if re.search(download_pattern, output):
                output = re.sub(download_pattern, "(document generation failed or skipped - please try again)", output)
                
        # Send text update if modified to replace hallucinated text on the client
        if output != original_output:
            yield {"type": "replace_text", "content": output}

        tool_data = {}
        for tool_name, tool_input, tool_output in completed_tools:
            tool_data[tool_name] = str(tool_output)[:500]

        if tool_data:
            yield {"type": "tool_result", "data": tool_data}

        # Step 6: Save conversation
        user_id = user.get("id", "") if user else ""
        save_text = f"📄 [Document: {doc_name}] {original_user_text}" if doc_name else text
        await save_message(session_id, user_id, "user", save_text, school_id)
        await save_message(session_id, user_id, "assistant", output, school_id, tool_data)

    except Exception as e:
        import traceback
        traceback.print_exc()
        error_msg = f"I encountered an error: {str(e)}. Please try again."
        yield {"type": "text", "content": error_msg}

    yield {"type": "done"}